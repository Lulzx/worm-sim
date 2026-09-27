#!/usr/bin/env python3
"""Independent source-array and directed-edge audit of native molecular import.

Does not establish biological validity of the declared class-to-cell transfer.
"""
import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import re
import h5py
import numpy as np
import openpyxl


def load(path):
    return json.loads(Path(path).read_text())


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--run', required=True)
    p.add_argument('--output', required=True)
    p.add_argument('--graph', default='runs/c302-audit.json')
    p.add_argument('--catalog', default='data/fenyves-transmitter-receptors.json')
    p.add_argument('--mapping', default='data/cengen-cell-map.json')
    p.add_argument('--xlsx', default='runs/molecular-source/journal.pcbi.1007974.s003.xlsx')
    p.add_argument('--hdf5', default='runs/molecular-source/cengen.h5')
    a = p.parse_args()
    run = Path(a.run)
    catalog, mapping, graph = load(a.catalog), load(a.mapping), load(a.graph)
    expression, evidence, report = [load(run/f'{name}.json') for name in ['expression', 'evidence', 'report']]
    assert digest(a.xlsx) == catalog['source']['sha256'] == '85959066fd7cbdbc2024d0ebb323b71c4365f4083bc85e555ee973f470697c47'
    assert digest(a.hdf5) == expression['source']['sha256'] == 'd8e6f6f2a25e05211676cfd9c3dd18b8a8bb5c3db4d673fd87b61ed91637ed23'
    workbook = openpyxl.load_workbook(a.xlsx, read_only=True, data_only=False)
    nt_rows = []
    for row in workbook['1. NT expr'].iter_rows(min_row=2, max_col=3):
        if row[0].value is None:
            continue
        assert all(c.data_type != 'f' for c in row)
        raw = row[0].value
        match = re.fullmatch(r'([A-Z]+)(\d+)', raw)
        canonical = match[1] + str(int(match[2])) if match else raw
        nt_rows.append(dict(neuron=canonical, source_neuron=raw, dominant=row[1].value, alternative=row[2].value, source_row=row[0].row))
    assert nt_rows == catalog['transmitters']
    receptors = []
    for col, (nt, sign) in enumerate([(n, s) for n in ['Glu','ACh','GABA'] for s in ['excitatory','inhibitory']], 1):
        for row in workbook['2. Receptor gene table'].iter_rows(min_row=2, min_col=col, max_col=col):
            cell = row[0]
            if cell.value is not None:
                assert cell.data_type != 'f'
                receptors.append(dict(gene=cell.value, transmitter=nt, polarity=sign, source_cell=cell.coordinate))
    assert receptors == catalog['receptors']
    requested = {r['gene'] for r in receptors}
    with h5py.File(a.hdf5) as f:
        th = expression['threshold']
        classes = f['neuron_ids'].asstr()[:].tolist()
        names = f[f'gene_names_th{th}'].asstr()[:].tolist()
        ids = f[f'gene_wbids_th{th}'].asstr()[:].tolist()
        source_matrix = f[f'tpm_th{th}'][:]
    assert expression['classes'] == classes
    present = sorted(requested.intersection(names))
    missing = sorted(requested.difference(names))
    assert expression['missing_genes'] == missing
    assert expression['genes'] == [dict(name=n, wormbase_id=ids[names.index(n)]) for n in present]
    native = np.asarray(expression['tpm'], dtype=np.float32).reshape(len(classes),len(present))
    np.testing.assert_array_equal(native,source_matrix[:,[names.index(n) for n in present]])
    cells = {n['id'] for n in graph['neurons']}
    assert set(mapping['cell_to_class']).isdisjoint(mapping['unmapped'])
    assert set(mapping['cell_to_class']) | set(mapping['unmapped']) == cells
    assert set(mapping['unmapped']) == {'AWCL','AWCR'}
    for neuron, class_ in {'DA9':'DA9','DB1':'DB01','VA12':'VA12','VB1':'VB01','VB2':'VB02','VC4':'VC_4_5','VC5':'VC_4_5','SABD':'SAB','URYVR':'URY'}.items():
        assert mapping['cell_to_class'][neuron] == class_
    nt = {r['neuron']: {r[k] for k in ['dominant','alternative'] if r[k] is not None} for r in nt_rows}
    expected = []
    for edge in sorted(graph['chemical'],key=lambda e:(e['pre'],e['post'])):
        pre, post = edge['pre'], edge['post']
        transmitters = nt[pre]
        signs = {'excitatory':set(), 'inhibitory':set()}
        unknown = set()
        if not transmitters:
            state = 'no_transmitter_evidence'
        elif post not in mapping['cell_to_class']:
            state = 'unmapped_postsynaptic_class'
        else:
            row = classes.index(mapping['cell_to_class'][post])
            for r in receptors:
                if r['transmitter'] not in transmitters:
                    continue
                if r['gene'] not in names:
                    unknown.add(r['gene'])
                elif source_matrix[row,names.index(r['gene'])] > 0:
                    signs[r['polarity']].add(r['gene'])
            if all(signs.values()): state='conflicting'
            elif unknown: state='incomplete_receptors'
            elif signs['excitatory']: state='excitatory'
            elif signs['inhibitory']: state='inhibitory'
            else: state='no_detected_receptor'
        expected.append(dict(pre=pre,post=post,transmitters=sorted(transmitters),expressed_excitatory=sorted(signs['excitatory']),expressed_inhibitory=sorted(signs['inhibitory']),missing_receptor_genes=sorted(unknown),state=state))
    assert evidence['edges'] == expected
    counts = dict(Counter(e['state'] for e in expected))
    assert counts == report['edge_states']
    receipt = dict(schema_version=1, native_report=report, source_files={'xlsx_sha256':digest(a.xlsx),'hdf5_sha256':digest(a.hdf5)},
                   artifacts={str(p):digest(p) for p in [Path(a.catalog),Path(a.mapping),Path(a.graph),run/'expression.json',run/'evidence.json']},
                   audit_script_sha256=digest(__file__),exact_tpm_values_checked=int(native.size),directed_edges_checked=len(expected),
                   limitations='Exact source-array, catalog-row, and qualitative-rule parity. Cell mapping is an explicit modeling transfer, not independently measured expression in anatomical cells. Zero TPM is thresholded non-detection, missing names remain unknown, and expression does not prove physiological synaptic polarity. No calibrated sign probability or fitted benchmark improvement claimed.')
    Path(a.output).write_text(json.dumps(receipt,indent=2)+'\n')
    print(json.dumps({'threshold':th,'edge_states':counts,'exact_tpm_values_checked':int(native.size)},indent=2))


if __name__ == '__main__':
    main()
