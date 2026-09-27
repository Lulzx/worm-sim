#!/usr/bin/env python3
"""Fetch pinned molecular sources and extract literal Fenyves table values.

Requires openpyxl only for read-only XLSX extraction. No upstream code is executed.
CeNGEN HDF5 loading and polarity inference are implemented by the Rust core.
"""
import hashlib
import json
from pathlib import Path
import re
from urllib.request import urlopen
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
REVISION = 'b2e13d88b670efcb3438aeacba2ad4bd6c383933'
BASE = f'https://raw.githubusercontent.com/francescorandi/wormneuroatlas/{REVISION}/wormneuroatlas/data/'
FILES = {
    'journal.pcbi.1007974.s003.xlsx': '85959066fd7cbdbc2024d0ebb323b71c4365f4083bc85e555ee973f470697c47',
    'cengen.h5': 'd8e6f6f2a25e05211676cfd9c3dd18b8a8bb5c3db4d673fd87b61ed91637ed23',
}


def main():
    directory = ROOT / 'runs/molecular-source'
    directory.mkdir(parents=True, exist_ok=True)
    for name, expected in FILES.items():
        path = directory / name
        if not path.exists():
            path.write_bytes(urlopen(BASE + name, timeout=60).read())
        actual = hashlib.sha256(path.read_bytes()).hexdigest()
        if actual != expected:
            raise ValueError(f'{name}: source hash mismatch')
    path = directory / 'journal.pcbi.1007974.s003.xlsx'
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    transmitters = []
    sheet = workbook['1. NT expr']
    assert [sheet.cell(1, i).value for i in range(1, 4)] == ['Class member', 'Dominant NT', 'Alternative NT']
    for row in sheet.iter_rows(min_row=2, max_col=3):
        name, dominant, alternative = [cell.value for cell in row]
        if name is None:
            assert dominant is None and alternative is None
            continue
        assert all(cell.data_type != 'f' for cell in row)
        # Only numeric zero padding changes. AWC left/right are retained literally.
        canonical = re.sub(r'0+(\d+)$', lambda m: str(int(m.group())), name)
        assert dominant in [None, 'Glu', 'ACh', 'GABA']
        assert alternative in [None, 'Glu', 'ACh', 'GABA']
        transmitters.append({'neuron': canonical, 'source_neuron': name,
                             'dominant': dominant, 'alternative': alternative,
                             'source_row': row[0].row})
    assert len(transmitters) == len({r['neuron'] for r in transmitters}) == 302
    canonical_ids = set(json.loads((ROOT/'data/c302-neuron-ids.json').read_text()))
    assert {r['neuron'] for r in transmitters} == canonical_ids
    sheet = workbook['2. Receptor gene table']
    receptors = []
    for col, (nt, polarity) in enumerate([(n, p) for n in ['Glu', 'ACh', 'GABA'] for p in ['excitatory', 'inhibitory']], 1):
        assert sheet.cell(1, col).value == nt + (' Pos' if polarity == 'excitatory' else ' Neg')
        for row in sheet.iter_rows(min_row=2, min_col=col, max_col=col):
            cell = row[0]
            if cell.value is not None:
                assert cell.data_type != 'f' and isinstance(cell.value, str)
                receptors.append({'gene': cell.value, 'transmitter': nt, 'polarity': polarity, 'source_cell': cell.coordinate})
    assert len(receptors) == len({r['gene'] for r in receptors}) == 62
    catalog = {'schema_version': 1,
               'source': {'url': BASE + path.name, 'sha256': FILES[path.name],
                          'version': 'Fenyves et al. 2020; doi:10.1371/journal.pcbi.1007974.s003',
                          'license': 'CC BY 4.0; https://creativecommons.org/licenses/by/4.0/'},
               'transmitters': transmitters, 'receptors': receptors}
    out = ROOT/'data/fenyves-transmitter-receptors.json'
    out.write_text(json.dumps(catalog, indent=2) + '\n')
    print(f'{len(transmitters)} cells, {len(receptors)} receptor genes; {out.relative_to(ROOT)}')


if __name__ == '__main__':
    main()
